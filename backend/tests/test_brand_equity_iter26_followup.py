"""Iteration 26 — Brand Equity follow-up (list-page scores, wave significance
via API, thesis export numbers, cache invalidation, compare-wave selector)."""
import os
import io
import time
import requests
import pytest
from dotenv import dotenv_values

fe = dotenv_values("/app/frontend/.env")
BASE_URL = (os.environ.get("REACT_APP_BACKEND_URL") or fe.get("REACT_APP_BACKEND_URL")).rstrip("/")

MALI_STUDY = "c5b3458f-b482-4859-ae15-b8c725e78d1c"
WAVE_STUDY = "zz-wavesig-0001"


# ---------- auth ----------
@pytest.fixture(scope="module")
def admin_headers():
    r = requests.post(f"{BASE_URL}/api/auth/login",
                      json={"email": "admin@swu.ac.th", "password": "admin123"})
    assert r.status_code == 200, r.text[:300]
    tok = r.json().get("access_token") or r.json().get("token")
    assert tok
    return {"Authorization": f"Bearer {tok}"}


# ---------- GET /studies/scores ----------
class TestListScoresEndpoint:
    def test_scores_endpoint_shape(self, admin_headers):
        r = requests.get(f"{BASE_URL}/api/brand-equity/studies/scores", headers=admin_headers)
        assert r.status_code == 200, r.text[:400]
        body = r.json()
        # accept either {scores: {...}} or a top-level map
        scores = body.get("scores") if isinstance(body, dict) and "scores" in body else body
        assert isinstance(scores, dict) and scores, f"empty scores: {body}"
        # find Mali + Wave rows
        assert MALI_STUDY in scores, f"Mali study missing in scores map (keys sample: {list(scores)[:3]})"
        row = scores[MALI_STUDY]
        for k in ("overall", "awareness", "quality", "identity", "loyalty",
                  "response_count", "locked"):
            assert k in row, f"key {k} missing in row for Mali: {row}"
        assert isinstance(row["locked"], bool)
        # Mali has >= 21 responses -> unlocked
        assert row["locked"] is False, row
        assert isinstance(row["overall"], (int, float))
        assert 0 <= row["overall"] <= 100

    def test_scores_endpoint_fast_on_second_call(self, admin_headers):
        # warm
        requests.get(f"{BASE_URL}/api/brand-equity/studies/scores", headers=admin_headers)
        t0 = time.perf_counter()
        r = requests.get(f"{BASE_URL}/api/brand-equity/studies/scores", headers=admin_headers)
        dt = time.perf_counter() - t0
        assert r.status_code == 200
        # cache should keep this well under 3s
        assert dt < 3.0, f"second call too slow: {dt:.2f}s"

    def test_scores_match_dashboard(self, admin_headers):
        r = requests.get(f"{BASE_URL}/api/brand-equity/studies/scores", headers=admin_headers)
        assert r.status_code == 200
        body = r.json()
        scores = body.get("scores") if isinstance(body, dict) and "scores" in body else body
        row = scores[MALI_STUDY]
        d = requests.get(f"{BASE_URL}/api/brand-equity/studies/{MALI_STUDY}/dashboard",
                         headers=admin_headers)
        assert d.status_code == 200
        agg = d.json().get("aggregate") or {}
        comp = agg.get("composite") or {}
        for k in ("overall", "awareness", "quality", "identity", "loyalty"):
            dash_val = comp.get(k)
            if dash_val is None:
                continue
            assert abs(row[k] - dash_val) < 0.6, \
                f"{k} mismatch: list={row[k]} dashboard={dash_val}"


# ---------- weight change -> cache invalidation ----------
class TestScoresCacheInvalidation:
    def test_weight_change_updates_scores_endpoint(self, admin_headers):
        # baseline
        r0 = requests.get(f"{BASE_URL}/api/brand-equity/studies/scores", headers=admin_headers)
        base = (r0.json().get("scores") or r0.json())[WAVE_STUDY]["overall"]

        # change weights heavily so BES has to move (loyalty dominates)
        put = requests.put(
            f"{BASE_URL}/api/brand-equity/studies/{WAVE_STUDY}",
            headers=admin_headers,
            json={"weights": {"awareness": 10, "quality": 10, "identity": 10, "loyalty": 70}},
        )
        assert put.status_code in (200, 204), put.text[:300]

        try:
            r1 = requests.get(f"{BASE_URL}/api/brand-equity/studies/scores", headers=admin_headers)
            new = (r1.json().get("scores") or r1.json())[WAVE_STUDY]["overall"]
            assert abs(new - base) > 0.5, f"scores did not update after weight change: {base} -> {new}"
        finally:
            # ALWAYS reset weights
            requests.put(
                f"{BASE_URL}/api/brand-equity/studies/{WAVE_STUDY}",
                headers=admin_headers,
                json={"weights": {"awareness": 25, "quality": 25, "identity": 25, "loyalty": 25}},
            )


# ---------- compare endpoint (wave selector + significance) ----------
class TestCompareWaveSelector:
    def test_compare_honours_wave_selector_and_returns_significance(self, admin_headers):
        r = requests.post(
            f"{BASE_URL}/api/brand-equity/studies/{WAVE_STUDY}/compare",
            headers=admin_headers,
            json={"wave_a": {"wave": "1"}, "wave_b": {"wave": "2"}},
        )
        assert r.status_code == 200, r.text[:600]
        body = r.json()
        sig = body.get("significance") or (body.get("compare") or {}).get("significance")
        assert sig, f"no significance block: keys={list(body.keys())}"
        by_dim = sig.get("by_dimension") or {}
        dims = {"awareness", "quality", "identity", "loyalty"}
        assert dims.issubset(set(by_dim.keys())), f"missing dims: {list(by_dim)}"
        # seeded so all 4 dimensions are significant, large effects
        for d in dims:
            row = by_dim[d]
            assert row.get("testable") is True, f"{d} not testable: {row}"
            assert row.get("significant") is True, f"{d} not significant: {row}"

    def test_compare_does_not_crash_on_measured_metadata(self, admin_headers):
        # Mali doesn't have two waves but the endpoint must not TypeError
        r = requests.post(
            f"{BASE_URL}/api/brand-equity/studies/{MALI_STUDY}/compare",
            headers=admin_headers,
            json={"wave_a": {"wave": "1"}, "wave_b": {"wave": "latest"}},
        )
        assert r.status_code in (200, 400, 422), r.text[:400]
        if r.status_code == 200:
            body = r.json()
            sig = body.get("significance") or (body.get("compare") or {}).get("significance")
            if sig:
                # If it returns something, tiny-wave dims must be non-testable
                for d, row in (sig.get("by_dimension") or {}).items():
                    if not row.get("testable"):
                        assert "reason" in row, f"{d} not-testable missing reason: {row}"
                        assert row.get("p_value") in (None, 0, 0.0, "") or "p_value" not in row \
                               or not row.get("significant"), \
                               f"{d} not testable but returned significant/p: {row}"


# ---------- thesis export numbers ----------
class TestThesisExportNumbers:
    def test_thesis_docx_matches_current_dashboard(self, admin_headers):
        # dashboard first (source of truth)
        d = requests.get(f"{BASE_URL}/api/brand-equity/studies/{MALI_STUDY}/dashboard",
                         headers=admin_headers)
        assert d.status_code == 200
        comp = ((d.json() or {}).get("aggregate") or {}).get("composite") or {}
        expected = {k: comp.get(k) for k in ("overall", "awareness", "quality", "identity", "loyalty")
                    if comp.get(k) is not None}
        assert expected

        r = requests.post(
            f"{BASE_URL}/api/brand-equity/studies/{MALI_STUDY}/thesis-export",
            headers=admin_headers,
        )
        assert r.status_code == 200, r.text[:400]
        assert len(r.content) > 5000, "docx too small"

        # parse docx
        try:
            from docx import Document
        except Exception:
            pytest.skip("python-docx not available")
        doc = Document(io.BytesIO(r.content))
        text_blob = "\n".join(p.text for p in doc.paragraphs)
        for tbl in doc.tables:
            for row in tbl.rows:
                for cell in row.cells:
                    text_blob += "\n" + cell.text

        # each expected score should appear rounded to 1 decimal in the docx text
        missing = []
        for k, v in expected.items():
            token1 = f"{v:.1f}"
            token2 = f"{round(v)}"
            if token1 not in text_blob and token2 not in text_blob:
                missing.append((k, v))
        assert not missing, f"docx missing dashboard values: {missing}"
