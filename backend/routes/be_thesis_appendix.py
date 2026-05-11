"""
be_thesis_appendix.py — Append a comprehensive visual & data appendix to an
AI-generated thesis DOCX so the final document includes EVERY dashboard chart
and full per-question statistics, not just the AI narrative.

Public API:
    enrich_thesis_docx(docx_bytes, study, aggregate) -> bytes

The appendix is injected at the end (after AI chapters + References):
  · ภาคผนวก ก — ภาพประกอบจาก Dashboard (TOM cloud, radars, perceptual map, identity-image, NPS, …)
  · ภาคผนวก ข — ตารางสถิติฉบับเต็มทุกคำถาม
  · ภาคผนวก ค — ตารางความถี่ open-ended (TOM, Recall, Aided)
  · ภาคผนวก ง — Identity vs Image trait-by-trait status
  · ภาคผนวก จ — Aaker 4 Perspectives composite breakdown
  · ภาคผนวก ฉ — Open-end Themes & Sample Quotes (Quality + Loyalty)
  · ภาคผนวก ช — Composite Summary recap
"""
import io
from typing import Any, Dict, List, Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt, RGBColor

from .be_thesis_charts import collect_all_charts

FONT = "TH SarabunPSK"
INK   = RGBColor(0x14, 0x16, 0x1A)
GREY  = RGBColor(0x5B, 0x5E, 0x63)
RED   = RGBColor(0xC8, 0x10, 0x2E)


def _set_run(run, *, size: float = 16, bold: bool = False,
             italic: bool = False, color: Optional[RGBColor] = None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color


def _h_section(doc: Document, text: str, *, level: int = 1):
    p = doc.add_paragraph()
    p.alignment = (WD_ALIGN_PARAGRAPH.CENTER if level == 0
                   else WD_ALIGN_PARAGRAPH.LEFT)
    r = p.add_run(text)
    _set_run(r, size=20 if level == 0 else 17, bold=True, color=INK)
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(8)


def _h_sub(doc: Document, text: str):
    p = doc.add_paragraph()
    r = p.add_run(text)
    _set_run(r, size=14, bold=True, color=RED)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)


def _para(doc: Document, text: str, *, italic: bool = False, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    _set_run(r, italic=italic, color=color or GREY, size=13)


def _add_image(doc: Document, png_buf: io.BytesIO, *,
               width_inches: float = 6.2, caption: Optional[str] = None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(png_buf, width=Inches(width_inches))
    if caption:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cp.add_run(caption)
        _set_run(r, size=12, italic=True, color=GREY)
        cp.paragraph_format.space_after = Pt(14)


def _data_table(
    doc: Document, headers: List[str], rows: List[List[Any]],
    *, font_size: float = 12,
):
    if not rows:
        return
    t = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        _set_run(r, size=font_size, bold=True, color=INK)
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            c = t.rows[ri].cells[ci]
            c.text = ""
            r = c.paragraphs[0].add_run("" if val is None else str(val))
            _set_run(r, size=font_size, color=INK)
    doc.add_paragraph()  # spacing


def _fmt(v, suffix: str = "") -> str:
    if v is None:
        return "—"
    if isinstance(v, float):
        return f"{round(v, 2)}{suffix}"
    return f"{v}{suffix}"


# ============================================================================

_QUESTION_KIND_LABEL_TH = {
    "unaided_top_of_mind":   "Top-of-Mind (เปิดให้พิมพ์ชื่อแรก)",
    "unaided_recall":        "Spontaneous Recall (เปิดให้พิมพ์เพิ่ม)",
    "aided_recognition":     "Aided Recognition (ติ๊กที่รู้จัก)",
    "logo_unaided":          "Logo Unaided",
    "logo_aided":            "Logo Aided",
    "color_recall":          "Color Recall",
    "brand_familiarity":     "Brand Familiarity (Likert-5)",
    "overall_quality":       "Overall Quality (Likert-7)",
    "attribute_reliability": "Quality · Reliability",
    "attribute_value":       "Quality · Value-for-Money",
    "attribute_design":      "Quality · Design",
    "brand_vs_competitor":   "Comparative vs Competitor",
    "image_traits_aided":    "Trait Pickup (BI-1, ติ๊ก 3-5)",
    "brand_promise_match":   "Brand Promise Match",
    "image_personality":     "Open Personality (AI normalized)",
    "image_one_word":        "Image · One-Word",
    "as_org_trust":          "As Org · Trust",
    "as_org_innovation":     "As Org · Innovation",
    "as_org_responsibility": "As Org · Responsibility",
    "as_person_personality": "As Person · Personality Grid",
    "as_person_relatable":   "As Person · Relatable",
    "as_person_aspirational":"As Person · Aspirational",
    "as_symbol_recognizability": "As Symbol · Recognizability",
    "as_symbol_distinctive": "As Symbol · Distinctive",
    "as_product_premium":    "As Product · Premium",
    "as_product_variety":    "As Product · Variety",
    "as_product_quality":    "As Product · Quality",
    "nps":                   "Net Promoter Score (0-10)",
    "repeat_intent":         "Repeat Intent (Likert-7)",
    "switching_resistance":  "Switching Resistance (Likert-7)",
    "reason_to_stay":        "Reason to stay (open-ended)",
    "quality_reasons":       "Quality reasons (open-ended)",
}


def _kind_label(k: Optional[str]) -> str:
    if not k:
        return "—"
    return _QUESTION_KIND_LABEL_TH.get(k, k)


# ============================================================================

def enrich_thesis_docx(
    docx_bytes: bytes,
    *, study: Dict[str, Any], aggregate: Dict[str, Any],
    mermaid_figures: Optional[List[tuple]] = None,
) -> bytes:
    """Open the AI-generated DOCX, append visual + data appendices, return new bytes.

    mermaid_figures: list of (block_id, png_bytes) from be_thesis_mermaid — rendered
                     before this function is called so the appendix can embed them.
    """
    doc = Document(io.BytesIO(docx_bytes))

    # Page break before appendix block so it doesn't collide with References
    doc.add_page_break()

    # =====================================================================
    # ภาคผนวก ก-mermaid — Conceptual / Framework diagrams from AI
    # =====================================================================
    if mermaid_figures:
        _h_section(doc, "ภาคผนวก ก-mermaid — แผนภาพ Conceptual / Framework", level=0)
        _para(doc,
              "แผนภาพ mermaid ที่ AI สร้างขึ้นในเนื้อหา (อาจอยู่ใน conceptual framework "
              "ของบทที่ 2 หรือ research methodology flowchart ของบทที่ 3) — "
              "ถูกเรนเดอร์เป็นภาพและฝังไว้ที่นี่.",
              italic=True)
        for bid, png_bytes in mermaid_figures:
            buf = io.BytesIO(png_bytes)
            _add_image(doc, buf, width_inches=6.2,
                       caption=f"ภาพ {bid} · แผนภาพแนวคิดจาก AI")
        doc.add_page_break()

    # =====================================================================
    # ภาคผนวก ก — ภาพประกอบจาก Dashboard
    # =====================================================================
    _h_section(doc, "ภาคผนวก ก — ภาพประกอบจาก Dashboard", level=0)
    _para(doc,
          "ส่วนนี้แสดงภาพประกอบทุกชิ้นจาก Brand Health Check Dashboard "
          "ที่สอดคล้องกับผลการวิเคราะห์ในบทที่ 4 — ทุกภาพคำนวณจากข้อมูลดิบจริงของกลุ่มตัวอย่าง.",
          italic=True)

    charts = collect_all_charts(study, aggregate)
    chart_meta = [
        ("bes_breakdown",       "ภาพ ก-1",  "Brand Equity Score — สรุปองค์ประกอบ 4 มิติ "
                                            "(Awareness · Quality · Identity · Loyalty)"),
        ("tom_cloud",          "ภาพ ก-2",  "Top-of-Mind Brand Cloud — แบรนด์เป้าหมาย "
                                            "(สีแดง) เทียบกับคู่แข่ง (สีเทา)"),
        ("unaided_recall_bars", "ภาพ ก-3", "Unaided Recall — แบรนด์อื่นที่ผู้ตอบนึกถึงเอง "
                                            "(ปกติแสดง 8-10 อันดับแรก)"),
        ("aided_recognition_bars", "ภาพ ก-4", "Aided Recognition — แบรนด์ที่ผู้ตอบรู้จัก "
                                            "เมื่อถูกแสดงรายชื่อให้เลือก"),
        ("familiarity_hist",   "ภาพ ก-5",  "Brand Familiarity Distribution (Likert-5)"),
        ("quality_radar",      "ภาพ ก-6",  "Quality Attributes Radar — ค่าเฉลี่ย Likert "
                                            "(1-7) ปรับเป็น 0-100"),
        ("overall_quality_hist","ภาพ ก-7", "Overall Quality Distribution (Likert-7)"),
        ("quality_reason_themes","ภาพ ก-8","Voice of Respondent — เหตุผลด้านคุณภาพ "
                                            "(Themes จาก open-end ที่ AI cluster)"),
        ("perspective_bars",   "ภาพ ก-9",  "Aaker 4 Perspectives — Brand Association "
                                            "(Org / Person / Symbol / Product)"),
        ("image_radar",        "ภาพ ก-10", "Multi-dimensional Image Radar — แบรนด์ vs คู่แข่ง"),
        ("perceptual_map",     "ภาพ ก-11", "Perceptual Map 2 มิติ — เลือกแกนที่แสดงความ "
                                            "ต่างของแบรนด์ในตลาด"),
        ("identity_image",     "ภาพ ก-12", "Identity (ที่แบรนด์ตั้งใจ) vs Image "
                                            "(ที่ผู้บริโภครับรู้)"),
        ("one_word_cloud",     "ภาพ ก-13", "One-word Identity Cloud — คำเดียวที่ผู้บริโภค "
                                            "ใช้บรรยายแบรนด์ (คำเขียวคือคำที่ตรงกับ Identity ตั้งใจ)"),
        ("nps_bar",            "ภาพ ก-14", "Net Promoter Score Distribution"),
        ("loyalty_means",      "ภาพ ก-15", "Loyalty Behaviour — Repeat Intent & "
                                            "Switching Resistance (mean ปรับเป็น 0-100)"),
        ("loyalty_themes",     "ภาพ ก-16", "Why Customers Stay — เหตุผลที่ลูกค้าอยู่ "
                                            "(Themes จาก open-end ที่ AI cluster)"),
    ]
    appended = 0
    for key, fig_label, caption in chart_meta:
        if key in charts:
            _add_image(doc, charts[key], width_inches=6.2,
                       caption=f"{fig_label} · {caption}")
            appended += 1
    if appended == 0:
        _para(doc, "ยังไม่มีภาพ — กรุณาตรวจสอบว่าข้อมูลเพียงพอ (เช่น n ≥ 10 + AI ทำงานเสร็จแล้ว)",
              italic=True, color=GREY)

    # =====================================================================
    # ภาคผนวก ข — ตารางสถิติฉบับเต็มทุกคำถาม
    # =====================================================================
    doc.add_page_break()
    _h_section(doc, "ภาคผนวก ข — ตารางสถิติฉบับเต็มทุกคำถาม", level=0)
    _para(doc,
          "ตารางนี้แสดงค่าสถิติ descriptive ของทุกคำถาม แยกตามมิติ Aaker — "
          "ค่าเฉลี่ย (mean), ส่วนเบี่ยงเบนมาตรฐาน (SD), จำนวนผู้ตอบ (n), "
          "Top-2-Box, Bottom-2-Box, Net Score, และคะแนนปรับ 0-100.",
          italic=True)

    questions = aggregate.get("questions") or []
    composite = aggregate.get("composite") or {}

    for dim_key, dim_label in [
        ("awareness", "ข.1  Brand Awareness — ค่าสถิติทุกคำถาม"),
        ("quality",   "ข.2  Perceived Quality — ค่าสถิติทุกคำถาม"),
        ("identity",  "ข.3  Brand Identity — ค่าสถิติทุกคำถาม"),
        ("loyalty",   "ข.4  Brand Loyalty — ค่าสถิติทุกคำถาม"),
    ]:
        _h_sub(doc, f"{dim_label} — Composite = {_fmt(composite.get(dim_key))}/100")
        dim_qs = [q for q in questions if q.get("dimension") == dim_key
                  and not q.get("compare_brand")]
        if not dim_qs:
            _para(doc, "ไม่มีคำถามในมิตินี้", italic=True)
            continue
        rows = []
        for q in dim_qs:
            s = q.get("stats") or {}
            rows.append([
                q.get("id", "—"),
                _kind_label(q.get("kind")),
                q.get("type", "—"),
                _fmt(s.get("n")),
                _fmt(s.get("mean")),
                _fmt(s.get("std")),
                _fmt(s.get("t2b"), "%"),
                _fmt(s.get("b2b"), "%"),
                _fmt(s.get("net")),
                _fmt(s.get("score_100")),
            ])
        _data_table(doc,
            ["ID", "Question kind", "Type", "n", "Mean", "SD",
             "T2B", "B2B", "Net", "Score(0-100)"],
            rows, font_size=11)

    # Comparative items in their own table
    comps = [q for q in questions if q.get("compare_brand")]
    if comps:
        _h_sub(doc, "ข.5  Comparative — vs each competitor")
        rows = []
        for q in comps:
            s = q.get("stats") or {}
            rows.append([
                q.get("compare_brand", "—"),
                q.get("axis_label") or _kind_label(q.get("kind")),
                _fmt(s.get("n")),
                _fmt(s.get("mean")),
                _fmt(s.get("favor_left_pct"), "%"),
                _fmt(s.get("favor_right_pct"), "%"),
                _fmt(s.get("score_100")),
            ])
        _data_table(doc,
            ["Competitor", "Axis", "n", "Mean", "Focal favors %",
             "Comp favors %", "Score (focal)"],
            rows, font_size=11)

    # =====================================================================
    # ภาคผนวก ค — ตารางความถี่ open-ended (TOM / Recall / Aided)
    # =====================================================================
    doc.add_page_break()
    _h_section(doc, "ภาคผนวก ค — ความถี่ของแบรนด์ที่ผู้ตอบกล่าวถึง", level=0)
    _para(doc,
          "ค่าความถี่ที่ AI ทำการ canonicalize "
          "(จัดกลุ่ม spelling variants ของแบรนด์เดียวกันให้เป็น bucket เดียว) — "
          "ใช้ในการคำนวณ Top-of-Mind, Spontaneous Recall, และ Aided Recognition.",
          italic=True)

    for kind, label in [
        ("unaided_top_of_mind", "ค.1  Top-of-Mind"),
        ("unaided_recall",      "ค.2  Spontaneous (Unaided) Recall"),
        ("aided_recognition",   "ค.3  Aided Recognition"),
    ]:
        q = next((qq for qq in questions if qq.get("kind") == kind), None)
        if not q:
            continue
        s = q.get("stats") or {}
        counts = s.get("counts") or []
        if not counts:
            continue
        _h_sub(doc, f"{label} (n = {s.get('n', 0)})")
        rows = []
        for c in counts[:25]:
            row = [
                c.get("label", "—"),
                _fmt(c.get("count")),
                _fmt(c.get("pct"), "%"),
            ]
            variants = c.get("variants") or []
            if isinstance(variants, list) and len(variants) > 1:
                row.append(", ".join(str(v) for v in variants[:8]))
            else:
                row.append("—")
            rows.append(row)
        _data_table(doc,
            ["Brand label", "Count", "% share", "AI-merged variants"],
            rows, font_size=11)

    # =====================================================================
    # ภาคผนวก ง — Identity vs Image trait-by-trait
    # =====================================================================
    ai_cache = study.get("ai_cache") or {}
    ii = ai_cache.get("identity_image") or {}
    intended = (study.get("identity") or {}).get("traits") or []
    aligned = ii.get("aligned") or []
    lost = ii.get("lost") or []
    unintended = ii.get("unintended") or []
    perceived_traits = ii.get("perceived_traits") or []

    if intended or perceived_traits:
        doc.add_page_break()
        _h_section(doc, "ภาคผนวก ง — Identity vs Image: ตารางเทียบรายลักษณะ", level=0)
        align = ii.get("alignment_score_0_100") or ii.get("alignment_score")
        _para(doc,
              f"คะแนน Identity-Image Alignment = {_fmt(align)}/100 — "
              "ตารางด้านล่างแสดงสถานะของแต่ละ trait: "
              "aligned (ตรง), lost (ตั้งใจแต่ไม่ถูกรับรู้), "
              "unintended (ถูกรับรู้แต่ไม่ได้ตั้งใจ).",
              italic=True)

        rows: List[List[Any]] = []
        for t in intended:
            status = ("aligned" if t in aligned
                      else "lost" if t in lost
                      else "lost")
            rows.append([t, "Intended", status, "—"])
        # Add perceived that are not in intended (unintended/aligned overlap already counted)
        seen_intended_low = {str(t).lower() for t in intended}
        for p in perceived_traits[:12]:
            if isinstance(p, dict):
                trait = p.get("trait") or p.get("label")
                count = p.get("count") or p.get("pct")
            else:
                trait = str(p); count = None
            if not trait or trait.lower() in seen_intended_low:
                continue
            status = ("aligned" if trait in aligned
                      else "unintended" if trait in unintended
                      else "perceived (other)")
            rows.append([trait, "Perceived", status, _fmt(count, "%" if isinstance(count, (int, float)) else "")])

        if rows:
            _data_table(doc,
                ["Trait", "Source", "Status", "Strength"],
                rows, font_size=11)

        if ii.get("narrative"):
            _h_sub(doc, "AI Narrative")
            _para(doc, ii["narrative"], color=INK)

    # =====================================================================
    # ภาคผนวก จ — Aaker 4 Perspectives composite breakdown
    # =====================================================================
    persp = (aggregate.get("association") or {}).get("by_perspective") or {}
    if any(v > 0 for v in persp.values()):
        doc.add_page_break()
        _h_section(doc, "ภาคผนวก จ — Aaker 4 Perspectives Breakdown", level=0)
        _para(doc,
              "Brand-as-Organization, Brand-as-Person, Brand-as-Symbol, Brand-as-Product — "
              "แต่ละมุมมองคำนวณจากค่าเฉลี่ย Likert-7 ของคำถามที่ tag ด้วย perspective นั้น.",
              italic=True)
        rows = [
            ["As Organization (องค์กร)",  _fmt(persp.get("as_organization"))],
            ["As Person (บุคลิก)",         _fmt(persp.get("as_person"))],
            ["As Symbol (สัญลักษณ์)",      _fmt(persp.get("as_symbol"))],
            ["As Product (ผลิตภัณฑ์)",     _fmt(persp.get("as_product"))],
        ]
        _data_table(doc,
            ["Perspective", "Score (0-100)"], rows, font_size=12)

        # Per-question contributions
        persp_qs = [q for q in questions if q.get("perspective")
                    and not q.get("compare_brand")]
        if persp_qs:
            _h_sub(doc, "จ.1  คำถามที่ feed เข้าแต่ละ Perspective")
            rows = []
            for q in persp_qs:
                s = q.get("stats") or {}
                rows.append([
                    q.get("perspective", "—"),
                    _kind_label(q.get("kind")),
                    _fmt(s.get("n")),
                    _fmt(s.get("mean")),
                    _fmt(s.get("score_100")),
                ])
            _data_table(doc,
                ["Perspective", "Question kind", "n", "Mean", "Score(0-100)"],
                rows, font_size=11)

    # =====================================================================
    # ภาคผนวก ฉ — Open-end Themes & Sample Quotes (Quality + Loyalty)
    # =====================================================================
    themes_map = (ai_cache.get("themes_by_question") or {})
    _OPEN_TYPES = {"open_text", "open_long", "text", "textarea", "open"}
    if themes_map:
        themes_appendix_added = False
        for dim_label, dim_key, kind_aliases in [
            ("ฉ.1  Quality — Open-end themes (เหตุผลด้านคุณภาพ)",
             "quality", ("quality_reasons", "quality_voice")),
            ("ฉ.2  Loyalty — Open-end themes (Why customers stay)",
             "loyalty", ("reason_to_stay", "loyalty_reasons")),
        ]:
            open_q = next((q for q in questions
                           if q.get("dimension") == dim_key
                           and (q.get("type") in _OPEN_TYPES
                                or q.get("kind") in kind_aliases)), None)
            if not open_q:
                continue
            themes_list = themes_map.get(open_q.get("id")) or []
            if not themes_list:
                continue
            if not themes_appendix_added:
                doc.add_page_break()
                _h_section(doc, "ภาคผนวก ฉ — Open-end Themes & Quotes "
                                "(เหตุผลที่ AI cluster จากคำตอบเปิด)", level=0)
                _para(doc,
                      "AI จัดกลุ่มคำตอบ open-end ออกเป็น 3-7 themes พร้อมยกตัวอย่างคำพูดจริง "
                      "(verbatim quotes) เพื่อให้ผู้อ่านงานวิจัยเข้าใจ texture ของข้อมูลเชิงคุณภาพ.",
                      italic=True)
                themes_appendix_added = True
            _h_sub(doc, dim_label)
            rows = []
            for t in themes_list:
                quote = ""
                quotes = t.get("sample_quotes") or []
                if quotes:
                    quote = str(quotes[0])[:160]
                rows.append([
                    str(t.get("name") or "—")[:60],
                    _fmt(t.get("count")),
                    quote or "—",
                ])
            _data_table(doc,
                ["Theme", "n", "Sample quote (verbatim)"],
                rows, font_size=11)

    # =====================================================================
    # Composite summary recap (last sanity table)
    # =====================================================================
    doc.add_page_break()
    _h_section(doc, "ภาคผนวก ช — สรุป Composite ทุกมิติ (Summary)", level=0)
    cm = aggregate.get("composite_meta") or {}
    rows = [
        ["BES (Overall Brand Equity)", _fmt(composite.get("overall")), "—", "—"],
    ]
    for d in ["awareness", "quality", "identity", "loyalty"]:
        meta = cm.get(d) or {}
        cov = meta.get("coverage")
        present = meta.get("present_kinds") or []
        missing = meta.get("missing_kinds") or []
        rows.append([
            d.capitalize(),
            _fmt(composite.get(d)),
            _fmt((cov or 0) * 100, "%") if cov is not None else "—",
            f"present: {len(present)} · missing: {len(missing)}",
        ])
    _data_table(doc,
        ["Dimension", "Score (0-100)", "Coverage", "Signal Coverage Detail"],
        rows, font_size=11)

    # Output
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
